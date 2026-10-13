import io
import shutil

import pymupdf
import pytest
from PIL import Image, ImageDraw, ImageFont

from conversor.extractores import extract


def _joined(fragments) -> str:
    return "\n\n".join(f.text for f in fragments)


def make_pdf(pages: list[str]) -> bytes:
    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        page.insert_text((72, 72), text, fontsize=12)
    return doc.tobytes()


def test_pdf_text_per_page():
    kind, fragments = extract(make_pdf(["Especificaciones técnicas", "Partida 2: Hormigón H-30"]), "eett.pdf")
    assert kind == "pdf"
    assert len(fragments) == 2
    assert "Hormigón H-30" in fragments[1].text
    assert fragments[1].metadata["page"] == 1
    assert fragments[1].metadata["page_label"] == "2"
    assert fragments[1].metadata["metodo"] == "texto"


@pytest.mark.skipif(not shutil.which("tesseract"), reason="requiere Tesseract")
def test_scanned_pdf_uses_ocr():
    image = Image.new("L", (1700, 600), 255)
    font = ImageFont.load_default(size=72)
    ImageDraw.Draw(image).text((60, 200), "CONTRATO DE OBRA 2026", fill=0, font=font)
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_image(page.rect, stream=buffer.getvalue())
    kind, fragments = extract(doc.tobytes(), "escaneado.pdf")

    assert fragments[0].metadata["metodo"] == "ocr"
    assert "CONTRATO" in fragments[0].text.upper()


def test_docx_headings_lists_and_tables():
    import docx

    document = docx.Document()
    document.add_heading("Contrato de construcción", level=1)
    document.add_paragraph("El mandante encarga la obra.")
    document.add_paragraph("Plazo de 180 días", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Partida", "Monto"
    table.cell(1, 0).text, table.cell(1, 1).text = "Excavación", "1.500.000"
    document.sections[0].footer.paragraphs[0].text = "Constructora — confidencial"
    buffer = io.BytesIO()
    document.save(buffer)

    kind, fragments = extract(buffer.getvalue(), "contrato.docx")
    text = _joined(fragments)
    assert kind == "docx"
    assert "## Contrato de construcción" in text
    assert "- Plazo de 180 días" in text
    assert "| Excavación | 1.500.000 |" in text
    assert "Pie de página: Constructora — confidencial" in text


def test_xlsx_sheets_as_tables():
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Presupuesto"
    sheet.append(["Partida", "Cantidad", "Precio unitario"])
    sheet.append(["Hormigón m3", 12.5, 98000])
    workbook.create_sheet("Vacía")
    buffer = io.BytesIO()
    workbook.save(buffer)

    kind, fragments = extract(buffer.getvalue(), "presupuesto.xlsx")
    assert kind == "xlsx"
    assert len(fragments) == 1
    assert fragments[0].metadata["hoja"] == "Presupuesto"
    assert "| Hormigón m3 | 12.5 | 98000 |" in fragments[0].text


def test_pptx_slides():
    import pptx

    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Avance de obra"
    slide.placeholders[1].text = "Obra gruesa 75%"
    buffer = io.BytesIO()
    presentation.save(buffer)

    kind, fragments = extract(buffer.getvalue(), "avance.pptx")
    assert kind == "pptx"
    assert fragments[0].text.startswith("## Diapositiva 1")
    assert "Obra gruesa 75%" in fragments[0].text


def test_detects_format_without_extension():
    # Open WebUI a veces no informa el tipo; se reconoce por el contenido.
    kind, fragments = extract(make_pdf(["Hola"]), "archivo_sin_extension")
    assert kind == "pdf"


def test_plain_text_windows_encoding():
    kind, fragments = extract("Año de construcción: 2026".encode("cp1252"), "nota.txt")
    assert kind == "text"
    assert fragments[0].text == "Año de construcción: 2026"
