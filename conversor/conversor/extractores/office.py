"""Word (.docx/.doc), Excel (.xlsx/.xls) y PowerPoint (.pptx) a Markdown."""

import datetime
import io
import shutil
import subprocess
import tempfile
from pathlib import Path

from .. import config
from .base import ExtractionError, Fragment, format_number, md_cell, md_table

# ---------------------------------------------------------------- Word .docx


def _docx_paragraph(paragraph) -> str:
    text = paragraph.text.strip()
    style = (paragraph.style.name if paragraph.style is not None else "") or ""
    if text:
        if style == "Title":
            text = f"# {text}"
        elif style.startswith("Heading"):
            level = style.removeprefix("Heading").strip()
            level = int(level) if level.isdigit() else 1
            text = f"{'#' * min(level + 1, 6)} {text}"
        elif "List" in style or (paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None):
            text = f"- {text}"

    # Cuadros de texto: python-docx no los incluye en paragraph.text.
    # Word guarda cada cuadro dos veces (versión moderna y VML), por eso se deduplica.
    lines = [text] if text else []
    for inner in paragraph._p.xpath(".//w:txbxContent//w:p"):
        box_text = "".join(t.text or "" for t in inner.xpath(".//w:t")).strip()
        if box_text and f"> {box_text}" not in lines:
            lines.append(f"> {box_text}")
    return "\n".join(lines)


def _docx_table(table) -> str:
    rows = []
    for row in table.rows:
        cells, previous = [], None
        for cell in row.cells:
            # Las celdas combinadas se repiten: se muestran una sola vez.
            if previous is not None and cell._tc is previous._tc:
                continue
            cells.append(md_cell(cell.text))
            previous = cell
        if any(cells):
            rows.append(cells)
    return md_table(rows)


def extract_docx(data: bytes) -> list[Fragment]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError(f"Documento Word dañado o ilegible: {exc}") from exc

    blocks = []
    headers, footers = [], []
    for section in document.sections:
        for part, target in ((section.header, headers), (section.footer, footers)):
            try:
                text = "\n".join(p.text.strip() for p in part.paragraphs if p.text.strip())
            except Exception:
                continue
            if text and text not in target:
                target.append(text)
    blocks += [f"Encabezado: {h}" for h in headers]

    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = _docx_paragraph(Paragraph(child, document))
        elif tag == "tbl":
            text = _docx_table(Table(child, document))
        elif tag == "sdt":  # controles de contenido (formularios)
            text = "\n".join(
                "".join(t.text or "" for t in p.xpath(".//w:t")) for p in child.xpath(".//w:p")
            ).strip()
        else:
            continue
        if text:
            blocks.append(text)

    blocks += [f"Pie de página: {f}" for f in footers]
    text = "\n\n".join(blocks).strip()
    return [Fragment(text)] if text else []


# ---------------------------------------------------------------- Word .doc (97-2003)


def extract_doc(data: bytes) -> list[Fragment]:
    if not shutil.which("antiword"):
        raise ExtractionError("Para leer .doc antiguos se necesita 'antiword'.")
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "documento.doc"
        path.write_bytes(data)
        proc = subprocess.run(
            ["antiword", "-m", "UTF-8.txt", "-w", "0", str(path)],
            capture_output=True,
            timeout=120,
        )
    if proc.returncode != 0:
        detail = proc.stderr.decode("utf-8", "replace").strip()
        raise ExtractionError(f"No se pudo leer el .doc (antiword): {detail}")
    text = proc.stdout.decode("utf-8", "replace").strip()
    return [Fragment(text)] if text else []


# ---------------------------------------------------------------- Excel


def _cell_value(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "VERDADERO" if value else "FALSO"
    if isinstance(value, float):
        return format_number(value)
    if isinstance(value, datetime.datetime):
        return value.strftime("%Y-%m-%d") if value.time() == datetime.time() else value.isoformat(" ")
    if isinstance(value, (datetime.date, datetime.time)):
        return value.isoformat()
    return md_cell(value)


def _sheet_fragment(name: str, index: int, rows_iter) -> Fragment | None:
    limit = config.settings.sheet_max_rows
    rows, truncated = [], False
    for values in rows_iter:
        cells = [_cell_value(v) for v in values]
        while cells and not cells[-1]:
            cells.pop()
        if not cells:
            continue
        if len(rows) >= limit:
            truncated = True
            break
        rows.append(cells)
    if not rows:
        return None
    text = f"## Hoja: {name}\n\n{md_table(rows)}"
    if truncated:
        text += f"\n\n(Hoja truncada: solo se incluyen las primeras {limit} filas con datos.)"
    return Fragment(text, {"page": index, "page_label": name, "hoja": name})


def extract_xlsx(data: bytes) -> list[Fragment]:
    import openpyxl

    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ExtractionError(f"Planilla Excel dañada o ilegible: {exc}") from exc
    try:
        fragments = []
        for index, sheet in enumerate(workbook.worksheets):
            fragment = _sheet_fragment(sheet.title, index, sheet.iter_rows(values_only=True))
            if fragment:
                fragments.append(fragment)
        return fragments
    finally:
        workbook.close()


def extract_xls(data: bytes) -> list[Fragment]:
    import xlrd

    try:
        workbook = xlrd.open_workbook(file_contents=data)
    except Exception as exc:
        raise ExtractionError(f"Planilla Excel (.xls) dañada o ilegible: {exc}") from exc

    def rows(sheet):
        for r in range(sheet.nrows):
            values = []
            for cell in sheet.row(r):
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        values.append(xlrd.xldate_as_datetime(cell.value, workbook.datemode))
                        continue
                    except Exception:
                        pass
                values.append(cell.value if cell.value != "" else None)
            yield values

    fragments = []
    for index, sheet in enumerate(workbook.sheets()):
        fragment = _sheet_fragment(sheet.name, index, rows(sheet))
        if fragment:
            fragments.append(fragment)
    return fragments


# ---------------------------------------------------------------- PowerPoint


def extract_pptx(data: bytes) -> list[Fragment]:
    import pptx

    try:
        presentation = pptx.Presentation(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError(f"Presentación dañada o ilegible: {exc}") from exc

    def shape_texts(shapes):
        for shape in shapes:
            if shape.shape_type == 6:  # grupo
                yield from shape_texts(shape.shapes)
            elif getattr(shape, "has_table", False) and shape.has_table:
                yield md_table([[md_cell(c.text) for c in row.cells] for row in shape.table.rows])
            elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    yield text

    fragments = []
    for index, slide in enumerate(presentation.slides):
        parts = list(shape_texts(slide.shapes))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"Notas del orador: {notes}")
        if parts:
            text = f"## Diapositiva {index + 1}\n\n" + "\n\n".join(parts)
            fragments.append(Fragment(text, {"page": index, "page_label": str(index + 1)}))
    return fragments
